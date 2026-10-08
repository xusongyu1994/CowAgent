"""Real PowerPoint containers must retain their text in both document tools."""

from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from threading import Thread

import pytest

from agent.tools.read.read import Read
from agent.tools.web_fetch.web_fetch import WebFetch


@pytest.mark.parametrize("tool_name", ["read", "web_fetch"])
@pytest.mark.parametrize("container", ["table", "merged_table", "group", "textbox", "empty"])
def test_powerpoint_container_text_is_read(tmp_path, monkeypatch, tool_name, container):
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    marker = "Quarterly revenue is 42 million"
    if container in ("table", "merged_table"):
        table = slide.shapes.add_table(2, 2, 0, 0, Inches(6), Inches(2)).table
        table.cell(0, 0).text = "Metric"
        table.cell(0, 1).text = "Result"
        table.cell(1, 0).text = marker
        table.cell(1, 1).text = "Confirmed"
        if container == "merged_table":
            table.cell(0, 0).merge(table.cell(0, 1))
            table.cell(0, 0).text = "Merged heading"
    elif container == "group":
        group = slide.shapes.add_group_shape()
        nested = group.shapes.add_group_shape()
        nested.shapes.add_textbox(0, 0, Inches(6), Inches(1)).text = marker
    elif container == "textbox":
        slide.shapes.add_textbox(0, 0, Inches(6), Inches(1)).text = marker
    if container != "empty":
        slide.shapes.add_textbox(0, Inches(3), Inches(6), Inches(1)).text = "Following text"
    path = tmp_path / "report.pptx"
    presentation.save(path)

    if tool_name == "read":
        result = Read({"cwd": str(tmp_path)}).execute({"path": str(path)})
        text = result.result.get("content", "")
    else:
        payload = path.read_bytes()

        class Handler(BaseHTTPRequestHandler):
            def do_GET(self):
                self.send_response(200)
                self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.presentationml.presentation")
                self.send_header("Content-Length", str(len(payload)))
                self.end_headers()
                self.wfile.write(payload)

            def log_message(self, *_args):
                pass

        monkeypatch.setenv("WEB_SECURITY_SSRF_PROTECTION", "false")
        server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        thread = Thread(target=server.serve_forever, daemon=True)
        thread.start()
        try:
            result = WebFetch({"cwd": str(tmp_path)}).execute({"url": f"http://127.0.0.1:{server.server_port}/report.pptx"})
        finally:
            server.shutdown()
            server.server_close()
            thread.join(timeout=5)
        text = result.result

    assert result.status == "success", result.result
    if container == "empty":
        assert marker not in text
    else:
        assert marker in text
        assert text.count(marker) == 1
        assert text.index(marker) < text.index("Following text")
    if container == "table":
        assert "Metric\tResult" in text
        assert f"{marker}\tConfirmed" in text

    if container == "merged_table":
        assert text.count("Merged heading") == 1
        assert text.index("Merged heading") < text.index(marker)
        assert f"{marker}\tConfirmed" in text
